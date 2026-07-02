# -*- coding: utf-8 -*-
"""
PetCare O2O 期末验收 PPT — 基于模板填充
策略：在模板副本上修改，保留背景图/版式/字体，替换占位文字 + 插入真实截图
作者：曾志洪  学号：2320160033
"""
import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Emu
from PIL import Image

# 路径
TEMPLATE = r"F:\trae code\petcare-o2o\.tmp_ppt_template.pptx"
OUT = r"F:\trae code\petcare-o2o\docs\PetCare O2O 期末验收报告.pptx"
SHOTS = r"F:\trae code\petcare-o2o\docs\build-and-deploy-guide\shots"
ESHOT = os.path.join(SHOTS, "report_shots")

ME = "曾志洪"
SID = "2320160033"
REPO = "https://github.com/zengmo20224/petcare-o2o"

# 复制模板为输出文件（保留全部背景图与版式）
shutil.copyfile(TEMPLATE, OUT)
prs = Presentation(OUT)
SW, SH = prs.slide_width, prs.slide_height  # EMU


# ---------- 工具函数 ----------
def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None

def first_text_shape(slide, key):
    """返回第一个文本含 key 的 shape"""
    for sh in slide.shapes:
        if sh.has_text_frame and key in sh.text_frame.text:
            return sh
    return None

def set_run_text(shape, para_idx, run_idx, text):
    """精准替换某段某run的文字，保留格式"""
    para = shape.text_frame.paragraphs[para_idx]
    if run_idx < len(para.runs):
        para.runs[run_idx].text = text

def replace_all_runs(shape, mapping):
    """按 (para_idx, run_idx) -> text 批量替换"""
    for (pi, ri), txt in mapping.items():
        set_run_text(shape, pi, ri, txt)

def clear_textframe(tf):
    """清空文本框所有段落，保留首段空"""
    tf.clear()

def add_text(tf, text, font_name="仿宋", size=None, color="FF0000", bold=True, align=None):
    """追加一段文字"""
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    p = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size) if size else Pt(24)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        from pptx.enum.text import PP_ALIGN
        p.alignment = align
    return p

def fit_image(slide, img_path, box_left_in, box_top_in, box_w_in, box_h_in):
    """在指定区域内等比缩放居中插入图片"""
    from PIL import Image as PILImage
    img = PILImage.open(img_path)
    iw, ih = img.size
    ratio = ih / iw
    # 限制宽高都不超出框
    w = box_w_in
    h = w * ratio
    if h > box_h_in:
        h = box_h_in
        w = h / ratio
    left = Inches(box_left_in + (box_w_in - w) / 2)
    top = Inches(box_top_in + (box_h_in - h) / 2)
    slide.shapes.add_picture(img_path, left, top, Inches(w), Inches(h))


slides = list(prs.slides)

# ============================================================
# 第1页 封面：曾志洪 / 学号 / 期末验收报告
# ============================================================
s1 = slides[0]
title1 = shape_by_name(s1, "标题 1")
# p0r0='xxx' p0r2=' ' p0r3='期末验收报告' p0r5='xxx  xxx xxx  xxxx'
# 重组：r0=PetCare O2O  r1=小组 r2=空格 r3=期末验收报告 r4=小组成员： r5=曾志洪 2320160033
replace_all_runs(title1, {
    (0, 0): "PetCare O2O",
    (0, 5): f"{ME}（{SID}）",
})

# ============================================================
# 第2页 主要内容（目录）— 模板已符合，仅微调最后一条"创新点、不足"
# ============================================================
s2 = slides[1]
box2 = shape_by_name(s2, "矩形 251909")
# 最后一段"创新点、不足" → 保持，内容贴切
# 无需改动

# ============================================================
# 第3页 项目概述 + H5/管理端截图
# ============================================================
s3 = slides[2]
t3 = first_text_shape(s3, "图书管理系统")  # 标题3 内容框
if t3:
    tf = t3.text_frame
    tf.clear()
    add_text(tf, "项目简介：", "仿宋", 26, "0000FF", True)
    add_text(tf, "PetCare O2O 是面向单体宠物门店的 O2O 平台，覆盖服务预约、商品零售、社区互动、营销活动四大业务域。", "仿宋", 22, "000000", False)
    add_text(tf, "技术架构：后端 Spring Boot 3.3 + MyBatis-Plus + MySQL 8；管理端 Vue3 + Element Plus；用户端 UniApp H5。", "仿宋", 22, "000000", False)
    add_text(tf, "交付成果：产品基线 v1.0.0-rc1 已建立，843 个测试全部通过，三端可访问。", "仿宋", 22, "000000", False)
    add_text(tf, "功能界面截图：", "仿宋", 24, "FF0000", True)
# 插入两张图：H5（竖图，左）+ 管理端（横图，右）
# 内容区约 left 1.7~12.6 in, top 4.0~7.0 in
fit_image(s3, os.path.join(SHOTS, "04_h5_home.png"), 2.0, 4.2, 2.6, 3.0)
fit_image(s3, os.path.join(SHOTS, "03_admin_dashboard.png"), 5.2, 4.3, 7.5, 2.7)

# ============================================================
# 第4页 小组分工及贡献（单人）
# ============================================================
s4 = slides[3]
# 该页只有标题，需添加内容文本框
tb4 = s4.shapes.add_textbox(Inches(1.7), Inches(1.6), Inches(10.5), Inches(5.2))
tf4 = tb4.text_frame
tf4.word_wrap = True
tf4.clear()
add_text(tf4, "本项目为单人课程作业，全部工作由曾志洪（学号 2320160033）独立完成。", "仿宋", 26, "0000FF", True)
add_text(tf4, "")
add_text(tf4, "主要承担角色与工作：", "仿宋", 24, "C00000", True)
add_text(tf4, "① 项目经理 / 配置管理工程师：编制配置管理计划，识别 76 个配置项，建立 8 个基线 tag。", "仿宋", 22, "000000", False)
add_text(tf4, "② 开发工程师：完成服务预约、商品订单、社区互动、营销活动四大模块三端开发。", "仿宋", 22, "000000", False)
add_text(tf4, "③ 测试工程师：编写 843 个单元 / 集成 / 契约测试，JaCoCo 类覆盖率 93%。", "仿宋", 22, "000000", False)
add_text(tf4, "④ DevOps 工程师：编写 Jenkins + GitHub Actions 流水线脚本，Docker 化一键部署。", "仿宋", 22, "000000", False)
add_text(tf4, "")
add_text(tf4, "文档贡献度：1.0（独立完成全部配置管理与持续集成工作）。", "仿宋", 24, "C00000", True)

# ============================================================
# 第5页 项目仓库目录结构 + git branch 图
# ============================================================
s5 = slides[4]
t5 = first_text_shape(s5, "界面截图")  # 标题3
if t5:
    tf = t5.text_frame
    tf.clear()
    add_text(tf, "版本库工具：Git + GitHub（主分支保护，PR 合并）", "仿宋", 24, "0000FF", True)
    add_text(tf, "仓库地址：" + REPO, "仿宋", 22, "000000", False)
# 目录树用独立文本框（等宽字体）
tb5 = s5.shapes.add_textbox(Inches(1.7), Inches(2.6), Inches(6.5), Inches(4.3))
tf5 = tb5.text_frame
tf5.word_wrap = True
tf5.clear()
tree_lines = [
    "petcare-o2o/",
    "├── src/                  后端源码(Java+测试)",
    "├── frontend/admin-web/   管理端(Vue3)",
    "├── frontend/miniapp/     用户端H5(UniApp)",
    "├── docs/                 配置管理文档",
    "├── nginx/                反向代理配置",
    "├── .github/workflows/    GitHub Actions",
    "├── Dockerfile            后端镜像",
    "├── docker-compose.yml    四服务编排",
    "├── Jenkinsfile           CI流水线",
    "├── schema.sql            数据库脚本",
    "└── pom.xml               Maven构建",
]
for i, ln in enumerate(tree_lines):
    from pptx.util import Pt
    p = tf5.paragraphs[0] if i == 0 else tf5.add_paragraph()
    run = p.add_run()
    run.text = ln
    run.font.name = "Consolas"
    run.font.size = Pt(16)
    run.font.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor.from_string("000000")
# 插入 git branch 图（右侧）
fit_image(s5, os.path.join(ESHOT, "e02_git_branch.png"), 8.5, 2.8, 4.3, 3.8)

# ============================================================
# 第6页 项目配置项清单
# ============================================================
s6 = slides[5]
t6 = first_text_shape(s6, "测试计划")  # 标题3
if t6:
    tf = t6.text_frame
    tf.clear()
    add_text(tf, "配置项分类清单（共 76 项，与配置管理计划一致）：", "仿宋", 24, "0000FF", True)
# 用表格呈现配置项分类
from pptx.util import Pt, Inches
rows, cols = 8, 3
tbl_left, tbl_top = Inches(1.7), Inches(2.4)
tbl_w, tbl_h = Inches(10.2), Inches(4.5)
table_shape = s6.shapes.add_table(rows, cols, tbl_left, tbl_top, tbl_w, tbl_h)
tbl = table_shape.table
tbl.columns[0].width = Inches(2.2)
tbl.columns[1].width = Inches(3.0)
tbl.columns[2].width = Inches(5.0)
ci_data = [
    ["类别码", "类别", "示例配置项"],
    ["SC", "源代码", "后端 Java / 前端 Vue-TS / 单元测试"],
    ["BS", "构建脚本", "pom.xml / package.json / Dockerfile"],
    ["CI", "CI/CD 流水线", "Jenkinsfile / GitHub Actions"],
    ["RC", "运行配置", "application*.yml / nginx.conf / .env"],
    ["DB", "数据库脚本", "schema.sql / data-dev.sql / migration"],
    ["DOC", "项目文档", "CMP / 构建指导书 / 部署指南 / 需求"],
    ["BL", "基线产物", "git tag / JaCoCo 报告 / Docker 镜像"],
]
from pptx.dml.color import RGBColor
for r, row_data in enumerate(ci_data):
    for c, val in enumerate(row_data):
        cell = tbl.cell(r, c)
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        run = para.runs[0]
        run.font.name = "仿宋"
        run.font.size = Pt(18)
        run.font.bold = (r == 0)
        run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else "000000")
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string("C00000" if r == 0 else "F2F2F2")

# ============================================================
# 第7页 成员分支开发、日志提交 + git log 图
# ============================================================
s7 = slides[6]
t7 = first_text_shape(s7, "各同学")  # 标题3
if t7:
    tf = t7.text_frame
    tf.clear()
    add_text(tf, "提交统计：累计 179 次提交，全部由曾志洪独立完成，遵循 Conventional Commits 规范。", "仿宋", 24, "0000FF", True)
    add_text(tf, "分支模型：main(发布基线) / develop(集成) / phase-2/7/10/11(里程碑) / hotfix", "仿宋", 22, "000000", False)
    add_text(tf, "提交日志与分支贡献情况截图：", "仿宋", 22, "FF0000", True)
# 插入 git log 图（左大）+ git branch（右小）
fit_image(s7, os.path.join(ESHOT, "e03_git_log.png"), 1.7, 4.0, 7.5, 3.0)
fit_image(s7, os.path.join(ESHOT, "e02_git_branch.png"), 9.4, 4.2, 3.4, 2.8)

# ============================================================
# 第8页 项目基线情况 + git tag 图
# ============================================================
s8 = slides[7]
t8 = first_text_shape(s8, "配置管理计划基线")  # 标题3
if t8:
    tf = t8.text_frame
    tf.clear()
    add_text(tf, "基线规划（符合 SemVer + 里程碑后缀命名规范）：", "仿宋", 24, "0000FF", True)
    add_text(tf, "共建立 8 个基线 tag：1 个功能基线 + 6 个里程碑基线 + 1 个产品基线。", "仿宋", 22, "000000", False)
    add_text(tf, "仓库实际基线情况截图：", "仿宋", 22, "FF0000", True)
# 基线表（左）+ git tag 图（右）
tbl8 = s8.shapes.add_table(6, 3, Inches(1.7), Inches(3.2), Inches(6.8), Inches(3.6)).table
tbl8.columns[0].width = Inches(2.2)
tbl8.columns[1].width = Inches(1.8)
tbl8.columns[2].width = Inches(2.8)
bl_data = [
    ["Tag", "类型", "内容"],
    ["v1.0.0-fb", "功能基线", "确立 V1 功能边界"],
    ["v1.0.0-m1", "里程碑", "H5 基础与演示数据"],
    ["v1.0.0-m4", "里程碑", "商品订单与库存金额"],
    ["v1.0.0-m6", "里程碑", "发布收口"],
    ["v1.0.0-rc1", "产品基线", "全部完成 + CI/CD 全通"],
]
for r, row_data in enumerate(bl_data):
    for c, val in enumerate(row_data):
        cell = tbl8.cell(r, c)
        cell.text = val
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.name = "仿宋"
        run.font.size = Pt(16)
        run.font.bold = (r == 0)
        run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else "000000")
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string("C00000" if r == 0 else "F2F2F2")
fit_image(s8, os.path.join(ESHOT, "e01_git_tag.png"), 8.8, 3.4, 4.0, 3.2)

# ============================================================
# 第9页 项目持续集成演示 + Jenkins 图 + docker ps 图
# ============================================================
s9 = slides[8]
t9 = first_text_shape(s9, "提交代买")  # 标题3
if t9:
    tf = t9.text_frame
    tf.clear()
    add_text(tf, "CI 流水线（Jenkins 主力 + GitHub Actions 补充）：", "仿宋", 24, "0000FF", True)
    add_text(tf, "编译 → 测试(843) → 打包 → Docker 构建 → 部署 → 健康检查 → 邮件通知，提交即自动触发。", "仿宋", 21, "000000", False)
    add_text(tf, "持续集成构建历史与部署运行证据：", "仿宋", 21, "FF0000", True)
fit_image(s9, os.path.join(ESHOT, "e04_jenkins.png"), 1.7, 3.9, 6.8, 3.2)
fit_image(s9, os.path.join(SHOTS, "02_docker_ps.png"), 8.7, 4.1, 4.2, 2.8)

# ============================================================
# 第10页 创新点、待改进、总结
# ============================================================
s10 = slides[9]
t10 = first_text_shape(s10, "可选择")  # 标题3
if t10:
    tf = t10.text_frame
    tf.clear()
    add_text(tf, "一、创新点 / 亮点", "仿宋", 24, "0000FF", True)
    add_text(tf, "1. 双 CI 互补：Jenkins 主力全流程，GitHub Actions 三 job 并行验证。", "仿宋", 20, "000000", False)
    add_text(tf, "2. Docker 多阶段构建：maven 编译 + JRE 运行，镜像更小，非 root 运行 + 健康检查。", "仿宋", 20, "000000", False)
    add_text(tf, "3. 一键部署健康关门：docker compose up --wait 一条命令完成构建/初始化/启动/健康检查。", "仿宋", 20, "000000", False)
    add_text(tf, "4. SCM 全闭环：8 基线 + 76 配置项 + 2 变更单 + 第 16 周审计证据，可审计可追溯。", "仿宋", 20, "000000", False)
    add_text(tf, "二、待改进", "仿宋", 24, "0000FF", True)
    add_text(tf, "1. 未引入 allure 测试报告（用 JaCoCo 覆盖率代替），可视化待加强。", "仿宋", 20, "000000", False)
    add_text(tf, "2. 前端覆盖率默认未启用，仅后端有 JaCoCo 覆盖率。", "仿宋", 20, "000000", False)
    add_text(tf, "三、总结", "仿宋", 24, "0000FF", True)
    add_text(tf, "独立完成从需求到产品基线的全流程 SCM 实践，掌握了 Git 分支策略、基线与变更控制、CI/CD 与容器化部署。", "仿宋", 20, "000000", False)

# 保存
prs.save(OUT)
print(f"OK -> {OUT}")
print(f"页数: {len(prs.slides)}")
